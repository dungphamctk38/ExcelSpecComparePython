from __future__ import annotations

import argparse
from collections import Counter
from dataclasses import dataclass, asdict
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


REPORT_COLUMNS = [
    "sheet",
    "change_type",
    "key",
    "cell",
    "column",
    "old_value",
    "new_value",
    "old_row",
    "new_row",
]

HEADER_FILL = PatternFill("solid", fgColor="4F81BD")
ADDED_FILL = PatternFill("solid", fgColor="C6EFCE")
REMOVED_FILL = PatternFill("solid", fgColor="FFC7CE")
CHANGED_FILL = PatternFill("solid", fgColor="FFEB9C")
WARNING_FILL = PatternFill("solid", fgColor="FCE4D6")

AUTO_KEY_COLUMNS = ["id", "spec id", "requirement id", "req id", "no", "no.", "number"]

STRUCTURE_ALIASES = [
    "structure",
    "struct",
    "struct name",
    "structure name",
    "\u69cb\u9020\u4f53",
    "\u69cb\u9020\u4f53\u540d",
]

FIELD_ALIASES = [
    "field",
    "field name",
    "member",
    "member name",
    "item",
    "item name",
    "\u9805\u76ee",
    "\u9805\u76ee\u540d",
    "\u30e1\u30f3\u30d0\u540d",
    "\u30e1\u30f3\u30d0\u30fc\u540d",
]


@dataclass
class Diff:
    sheet: str
    change_type: str
    key: str = ""
    cell: str = ""
    column: str = ""
    old_value: str = ""
    new_value: str = ""
    old_row: int | None = None
    new_row: int | None = None


def main() -> int:
    parser = argparse.ArgumentParser(description="Compare two Excel spec files and create a colored report.")
    parser.add_argument("old_file", help="Old/baseline Excel file.")
    parser.add_argument("new_file", help="New/current Excel file.")
    parser.add_argument("-o", "--output", default="diff_report.xlsx", help="Output .xlsx report path.")
    parser.add_argument("--mode", choices=["auto", "keyed", "cell"], default="auto")
    parser.add_argument("--profile", choices=["generic", "struct-spec"], default="struct-spec")
    parser.add_argument("--key-column", action="append", default=[], help="Key column. Can be used multiple times.")
    parser.add_argument("--header-row", type=int, default=1)
    parser.add_argument("--ignore-case", action="store_true")
    parser.add_argument("--no-trim", action="store_true")
    args = parser.parse_args()

    old_file = Path(args.old_file)
    new_file = Path(args.new_file)
    output = Path(args.output)

    if not old_file.exists():
        parser.error(f"Old file not found: {old_file}")
    if not new_file.exists():
        parser.error(f"New file not found: {new_file}")
    if args.header_row < 1:
        parser.error("--header-row must be 1 or greater.")

    diffs = compare_excel_files(
        old_file=old_file,
        new_file=new_file,
        mode=args.mode,
        profile=args.profile,
        key_columns=args.key_column,
        header_row=args.header_row,
        trim=not args.no_trim,
        ignore_case=args.ignore_case,
    )
    write_excel_report(diffs, output)

    print(f"Compared: {old_file} -> {new_file}")
    print(f"Report: {output}")
    print(f"Total differences: {len(diffs)}")
    for change_type, count in sorted(Counter(diff.change_type for diff in diffs).items()):
        print(f"  {change_type}: {count}")

    return 1 if diffs else 0


def compare_excel_files(
    old_file: Path,
    new_file: Path,
    mode: str,
    profile: str,
    key_columns: list[str],
    header_row: int,
    trim: bool,
    ignore_case: bool,
) -> list[Diff]:
    old_wb = load_workbook(old_file, data_only=True)
    new_wb = load_workbook(new_file, data_only=True)
    diffs: list[Diff] = []

    old_sheets = set(old_wb.sheetnames)
    new_sheets = set(new_wb.sheetnames)

    for sheet in sorted(old_sheets - new_sheets):
        diffs.append(Diff(sheet=sheet, change_type="sheet_removed"))
    for sheet in sorted(new_sheets - old_sheets):
        diffs.append(Diff(sheet=sheet, change_type="sheet_added"))

    for sheet in old_wb.sheetnames:
        if sheet not in new_sheets:
            continue

        old_ws = old_wb[sheet]
        new_ws = new_wb[sheet]
        keys = find_key_columns(old_ws, new_ws, mode, profile, key_columns, header_row)

        if keys:
            diffs.extend(compare_sheet_by_key(old_ws, new_ws, keys, header_row, trim, ignore_case))
        elif mode == "keyed" or key_columns:
            diffs.append(Diff(sheet=sheet, change_type="key_not_found", old_value="No key column found."))
        else:
            diffs.extend(compare_sheet_by_cell(old_ws, new_ws, trim, ignore_case))

    return diffs


def compare_sheet_by_key(old_ws, new_ws, key_columns: list[str], header_row: int, trim: bool, ignore_case: bool) -> list[Diff]:
    diffs: list[Diff] = []
    old_headers = read_headers(old_ws, header_row)
    new_headers = read_headers(new_ws, header_row)
    old_header_map = {normalize_header(name): col for col, name in old_headers.items()}
    new_header_map = {normalize_header(name): col for col, name in new_headers.items()}

    old_columns = set(old_header_map)
    new_columns = set(new_header_map)

    for column in sorted(old_columns - new_columns):
        diffs.append(Diff(sheet=old_ws.title, change_type="column_removed", column=old_headers[old_header_map[column]]))
    for column in sorted(new_columns - old_columns):
        diffs.append(Diff(sheet=old_ws.title, change_type="column_added", column=new_headers[new_header_map[column]]))

    old_rows = rows_by_key(old_ws, old_headers, key_columns, header_row)
    new_rows = rows_by_key(new_ws, new_headers, key_columns, header_row)
    diffs.extend(duplicate_key_diffs(old_ws.title, old_rows, "old"))
    diffs.extend(duplicate_key_diffs(new_ws.title, new_rows, "new"))

    old_unique_keys = {key for key, row_data in old_rows.items() if not row_data["duplicate_rows"]}
    new_unique_keys = {key for key, row_data in new_rows.items() if not row_data["duplicate_rows"]}

    for key in sorted(old_unique_keys - new_unique_keys):
        old_row_number = old_rows[key]["row_number"]
        old_row = old_rows[key]["row"]
        diffs.append(Diff(sheet=old_ws.title, change_type="row_removed", key=key, old_value=row_preview(old_row), old_row=old_row_number))

    for key in sorted(new_unique_keys - old_unique_keys):
        new_row_number = new_rows[key]["row_number"]
        new_row = new_rows[key]["row"]
        diffs.append(Diff(sheet=old_ws.title, change_type="row_added", key=key, new_value=row_preview(new_row), new_row=new_row_number))

    key_names = {normalize_header(name) for name in key_columns}
    for key in sorted(old_unique_keys & new_unique_keys):
        old_row_number = old_rows[key]["row_number"]
        old_row = old_rows[key]["row"]
        new_row_number = new_rows[key]["row_number"]
        new_row = new_rows[key]["row"]
        for column in sorted(old_columns & new_columns):
            if column in key_names:
                continue
            old_name = old_headers[old_header_map[column]]
            new_name = new_headers[new_header_map[column]]
            old_value = old_row.get(old_name)
            new_value = new_row.get(new_name)
            if values_equal(old_value, new_value, trim, ignore_case):
                continue
            diffs.append(
                Diff(
                    sheet=old_ws.title,
                    change_type="value_changed",
                    key=key,
                    column=old_name,
                    old_value=display(old_value),
                    new_value=display(new_value),
                    old_row=old_row_number,
                    new_row=new_row_number,
                )
            )

    return diffs


def compare_sheet_by_cell(old_ws, new_ws, trim: bool, ignore_case: bool) -> list[Diff]:
    diffs: list[Diff] = []
    for row in range(1, max(old_ws.max_row, new_ws.max_row) + 1):
        for col in range(1, max(old_ws.max_column, new_ws.max_column) + 1):
            old_value = old_ws.cell(row=row, column=col).value
            new_value = new_ws.cell(row=row, column=col).value
            if values_equal(old_value, new_value, trim, ignore_case):
                continue
            diffs.append(
                Diff(
                    sheet=old_ws.title,
                    change_type="cell_changed",
                    cell=f"{get_column_letter(col)}{row}",
                    old_value=display(old_value),
                    new_value=display(new_value),
                    old_row=row,
                    new_row=row,
                )
            )
    return diffs


def find_key_columns(old_ws, new_ws, mode: str, profile: str, key_columns: list[str], header_row: int) -> list[str]:
    if mode == "cell":
        return []

    old_names = available_headers(old_ws, header_row)
    new_names = available_headers(new_ws, header_row)

    if key_columns:
        return key_columns if all(normalize_header(col) in old_names and normalize_header(col) in new_names for col in key_columns) else []

    if profile == "struct-spec":
        structure = first_common_header(old_names, new_names, STRUCTURE_ALIASES)
        field = first_common_header(old_names, new_names, FIELD_ALIASES)
        if structure and field:
            return [structure, field]

    for candidate in AUTO_KEY_COLUMNS:
        if candidate in old_names and candidate in new_names:
            return [old_names[candidate]]

    return []


def read_headers(ws, header_row: int) -> dict[int, str]:
    headers = {}
    for col in range(1, ws.max_column + 1):
        value = ws.cell(row=header_row, column=col).value
        if value is not None and str(value).strip():
            headers[col] = str(value).strip()
    return headers


def available_headers(ws, header_row: int) -> dict[str, str]:
    return {normalize_header(name): name for name in read_headers(ws, header_row).values()}


def first_common_header(old_names: dict[str, str], new_names: dict[str, str], aliases: list[str]) -> str:
    for alias in aliases:
        normalized = normalize_header(alias)
        if normalized in old_names and normalized in new_names:
            return old_names[normalized]
    return ""


def rows_by_key(ws, headers: dict[int, str], key_columns: list[str], header_row: int) -> dict[str, dict[str, Any]]:
    header_map = {normalize_header(name): col for col, name in headers.items()}
    key_indexes = [header_map[normalize_header(name)] for name in key_columns]
    rows = {}

    for row_number in range(header_row + 1, ws.max_row + 1):
        row = {name: ws.cell(row=row_number, column=col).value for col, name in headers.items()}
        if all(value is None or str(value).strip() == "" for value in row.values()):
            continue
        key = " | ".join(display(ws.cell(row=row_number, column=col).value).strip() for col in key_indexes)
        key = key or f"blank_key_row_{row_number}"
        if key in rows:
            rows[key]["duplicate_rows"].append(row_number)
            continue
        rows[key] = {"row_number": row_number, "row": row, "duplicate_rows": []}

    return rows


def duplicate_key_diffs(sheet: str, rows: dict[str, dict[str, Any]], side: str) -> list[Diff]:
    diffs = []
    for key, row_data in rows.items():
        duplicate_rows = row_data["duplicate_rows"]
        if not duplicate_rows:
            continue
        all_rows = [row_data["row_number"], *duplicate_rows]
        diffs.append(
            Diff(
                sheet=sheet,
                change_type=f"duplicate_key_{side}",
                key=key,
                old_value=", ".join(str(row) for row in all_rows) if side == "old" else "",
                new_value=", ".join(str(row) for row in all_rows) if side == "new" else "",
                old_row=row_data["row_number"] if side == "old" else None,
                new_row=row_data["row_number"] if side == "new" else None,
            )
        )
    return diffs


def write_excel_report(diffs: list[Diff], output: Path) -> None:
    output.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    summary_ws = wb.active
    summary_ws.title = "Summary"
    detail_ws = wb.create_sheet("Differences")

    summary_ws.append(["change_type", "count"])
    for change_type, count in sorted(Counter(diff.change_type for diff in diffs).items()):
        summary_ws.append([change_type, count])
        apply_fill(summary_ws[summary_ws.max_row], fill_for(change_type))
    summary_ws.append(["total", len(diffs)])
    summary_ws.append([])
    summary_ws.append(["Color", "Meaning"])
    apply_fill(summary_ws[summary_ws.max_row], HEADER_FILL)
    for label, meaning, fill in [
        ("Green", "Added sheet, column, or row.", ADDED_FILL),
        ("Red", "Removed sheet, column, or row.", REMOVED_FILL),
        ("Yellow", "Changed value or cell.", CHANGED_FILL),
        ("Orange", "Warning.", WARNING_FILL),
    ]:
        summary_ws.append([label, meaning])
        apply_fill(summary_ws[summary_ws.max_row], fill)

    detail_ws.append(REPORT_COLUMNS)
    for diff in diffs:
        data = asdict(diff)
        detail_ws.append([data[column] for column in REPORT_COLUMNS])
        apply_fill(detail_ws[detail_ws.max_row], fill_for(diff.change_type))

    for ws in [summary_ws, detail_ws]:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = HEADER_FILL
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        auto_width(ws)

    wb.save(output)


def values_equal(old_value: Any, new_value: Any, trim: bool, ignore_case: bool) -> bool:
    return normalize_value(old_value, trim, ignore_case) == normalize_value(new_value, trim, ignore_case)


def normalize_value(value: Any, trim: bool, ignore_case: bool) -> str:
    text = display(value)
    if trim:
        text = " ".join(text.strip().split())
    if ignore_case:
        text = text.casefold()
    return text


def normalize_header(value: str) -> str:
    return " ".join(str(value).replace("_", " ").strip().casefold().split())


def display(value: Any) -> str:
    return "" if value is None else str(value)


def row_preview(row: dict[str, Any]) -> str:
    parts = [f"{key}={display(value)}" for key, value in row.items() if value not in (None, "")]
    return "; ".join(parts[:8])


def fill_for(change_type: str) -> PatternFill:
    if change_type.endswith("_added") or change_type == "row_added":
        return ADDED_FILL
    if change_type.endswith("_removed") or change_type == "row_removed":
        return REMOVED_FILL
    if change_type in {"value_changed", "cell_changed"}:
        return CHANGED_FILL
    if change_type == "key_not_found" or change_type.startswith("duplicate_key"):
        return WARNING_FILL
    return PatternFill()


def apply_fill(row, fill: PatternFill) -> None:
    for cell in row:
        cell.fill = fill


def auto_width(ws) -> None:
    for column_cells in ws.columns:
        column_letter = get_column_letter(column_cells[0].column)
        max_len = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column_cells)
        ws.column_dimensions[column_letter].width = min(max(max_len + 2, 12), 80)


if __name__ == "__main__":
    raise SystemExit(main())
