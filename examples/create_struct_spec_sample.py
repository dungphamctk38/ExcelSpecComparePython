from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


HEADERS = ["Structure Name", "Field Name", "Data Type", "Length", "Required", "Description"]


def _style_header(ws) -> None:
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    widths = [20, 20, 16, 10, 10, 36]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=index).column_letter].width = width
    ws.freeze_panes = "A2"


def _create_workbook(rows: list[list[str]]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Structure Definition"
    ws.append(HEADERS)
    for row in rows:
        ws.append(row)
    _style_header(ws)
    return wb


def save_old(path: Path) -> None:
    rows = [
        ["UserInfo", "user_id", "uint32", "4", "Y", "User ID"],
        ["UserInfo", "user_name", "char[32]", "32", "Y", "User name"],
        ["UserInfo", "email", "char[128]", "128", "N", "Email address"],
        ["OrderInfo", "order_id", "uint64", "8", "Y", "Order ID"],
        ["OrderInfo", "amount", "int32", "4", "Y", "Order amount"],
    ]
    _create_workbook(rows).save(path)


def save_new(path: Path) -> None:
    rows = [
        ["UserInfo", "user_id", "uint64", "8", "Y", "User ID"],
        ["UserInfo", "user_name", "char[64]", "64", "Y", "User name"],
        ["UserInfo", "phone_no", "char[20]", "20", "N", "Phone number"],
        ["OrderInfo", "order_id", "uint64", "8", "Y", "Order ID"],
        ["OrderInfo", "amount", "decimal", "8", "Y", "Order amount"],
    ]
    _create_workbook(rows).save(path)


if __name__ == "__main__":
    base_dir = Path(__file__).resolve().parent
    save_old(base_dir / "struct_spec_old.xlsx")
    save_new(base_dir / "struct_spec_new.xlsx")
    print("Created examples/struct_spec_old.xlsx")
    print("Created examples/struct_spec_new.xlsx")

